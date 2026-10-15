"""
Multi-Format Exporter for HumanizeAI.pro
Exports documents into Markdown (.md), Word (.docx), PDF (.pdf), and LaTeX (.tex).
"""

import os
from pathlib import Path
from typing import Optional


class DocumentFormatter:
    """Exports structured text to multiple target formats."""

    @staticmethod
    def export(text: str, output_path: str, format_type: str = "md") -> str:
        """Export text to specified file format.

        Args:
            text: Markdown or formatted text.
            output_path: Target file path.
            format_type: 'md', 'docx', 'pdf', or 'latex'.

        Returns:
            Absolute path to generated file.
        """
        fmt = format_type.lower().strip()
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        if fmt in ("md", "markdown", "txt"):
            out.write_text(text, encoding="utf-8")
            return str(out.resolve())

        elif fmt == "docx":
            return DocumentFormatter._export_docx(text, out)

        elif fmt == "pdf":
            return DocumentFormatter._export_pdf(text, out)

        elif fmt in ("latex", "tex"):
            return DocumentFormatter._export_latex(text, out)

        else:
            out.write_text(text, encoding="utf-8")
            return str(out.resolve())

    @staticmethod
    def _export_docx(text: str, out_path: Path) -> str:
        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = docx.Document()

            # Set margins to 1 inch
            for section in doc.sections:
                section.top_margin = Inches(1.0)
                section.bottom_margin = Inches(1.0)
                section.left_margin = Inches(1.0)
                section.right_margin = Inches(1.0)

            paragraphs = text.split("\n\n")
            for p_str in paragraphs:
                p_clean = p_str.strip()
                if not p_clean:
                    continue

                if p_clean.startswith("# "):
                    doc.add_heading(p_clean[2:], level=1)
                elif p_clean.startswith("## "):
                    doc.add_heading(p_clean[3:], level=2)
                elif p_clean.startswith("### "):
                    doc.add_heading(p_clean[4:], level=3)
                elif p_clean.startswith("- ") or p_clean.startswith("* "):
                    for line in p_clean.split("\n"):
                        item = line.lstrip("- *").strip()
                        if item:
                            doc.add_paragraph(item, style="List Bullet")
                else:
                    p = doc.add_paragraph(p_clean)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(6)

            doc.save(str(out_path))
            return str(out_path.resolve())
        except Exception as e:
            raise RuntimeError(f"DOCX generation failed: {e}")

    @staticmethod
    def _export_pdf(text: str, out_path: Path) -> str:
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib import colors

            doc = SimpleDocTemplate(
                str(out_path),
                pagesize=letter,
                rightMargin=72,
                leftMargin=72,
                topMargin=72,
                bottomMargin=72,
            )

            styles = getSampleStyleSheet()
            normal_style = styles["Normal"]
            normal_style.fontSize = 11
            normal_style.leading = 15
            normal_style.textColor = colors.HexColor("#222222")

            h1_style = styles["Heading1"]
            h1_style.fontSize = 18
            h1_style.leading = 22
            h1_style.textColor = colors.HexColor("#111827")

            h2_style = styles["Heading2"]
            h2_style.fontSize = 14
            h2_style.leading = 18
            h2_style.textColor = colors.HexColor("#1f2937")

            story = []
            for p_str in text.split("\n\n"):
                p_clean = p_str.strip()
                if not p_clean:
                    continue

                if p_clean.startswith("# "):
                    story.append(Paragraph(p_clean[2:], h1_style))
                    story.append(Spacer(1, 10))
                elif p_clean.startswith("## "):
                    story.append(Paragraph(p_clean[3:], h2_style))
                    story.append(Spacer(1, 8))
                elif p_clean.startswith("### "):
                    story.append(Paragraph(p_clean[4:], styles["Heading3"]))
                    story.append(Spacer(1, 6))
                else:
                    # Escape XML entities for ReportLab
                    safe_text = (
                        p_clean.replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;")
                    )
                    story.append(Paragraph(safe_text, normal_style))
                    story.append(Spacer(1, 10))

            doc.build(story)
            return str(out_path.resolve())
        except Exception as e:
            raise RuntimeError(f"PDF generation failed: {e}")

    @staticmethod
    def _export_latex(text: str, out_path: Path) -> str:
        # Generate clean academic article template
        body_lines = []
        for p in text.split("\n\n"):
            p_clean = p.strip()
            if not p_clean:
                continue
            if p_clean.startswith("# "):
                body_lines.append(f"\\section{{{p_clean[2:]}}}")
            elif p_clean.startswith("## "):
                body_lines.append(f"\\subsection{{{p_clean[3:]}}}")
            elif p_clean.startswith("### "):
                body_lines.append(f"\\subsubsection{{{p_clean[4:]}}}")
            else:
                body_lines.append(p_clean + "\n")

        latex_template = f"""\\documentclass[11pt,a4paper]{{article}}
\\usepackage[utf8]{{inputenc}}
\\usepackage[margin=1in]{{geometry}}
\\usepackage{{amsmath,amssymb}}
\\usepackage{{hyperref}}
\\usepackage{{setspace}}
\\onehalfspacing

\\title{{Humanized Document}}
\\author{{HumanizeAI Pro Swarm}}
\\date{{\\today}}

\\begin{{document}}
\\maketitle

{chr(10).join(body_lines)}

\\end{{document}}
"""
        out_path.write_text(latex_template, encoding="utf-8")
        return str(out_path.resolve())
